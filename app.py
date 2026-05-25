"""
JIKAP Manuscript Formatter - Streamlit Web App
Jurnal Informasi dan Komunikasi Administrasi Perkantoran
PAP FKIP Universitas Sebelas Maret
"""

import streamlit as st
import io
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─────────────────────────────────────────────
# CONSTANTS
# ─────────────────────────────────────────────
FONT_ARIAL = "Arial"
FONT_TNR   = "Times New Roman"
MARGIN_CM  = 3.0

SECTIONS_H1 = {
    "introduction", "pendahuluan",
    "research methods", "metode penelitian", "metodologi penelitian", "method",
    "results and discussion", "hasil dan pembahasan", "hasil penelitian dan pembahasan",
    "conclusion", "kesimpulan", "conclusions",
    "references", "daftar pustaka",
    "acknowledgments", "acknowledgement", "ucapan terima kasih",
}

SECTIONS_H2 = {
    "research results", "hasil penelitian",
    "discussion", "pembahasan",
}

# ─────────────────────────────────────────────
# REFERENCE PARSER & APA 7 FORMATTER
# ─────────────────────────────────────────────

def clean_text(t):
    return re.sub(r'\s+', ' ', t).strip()

def parse_author_string(raw):
    raw = raw.strip()
    raw = re.sub(r'\s*&\s*$', '', raw).strip()
    parts = re.split(r',\s*&\s*|\s*&\s*|;\s*', raw)
    authors = []
    i = 0
    while i < len(parts):
        part = parts[i].strip()
        if not part:
            i += 1; continue
        if i + 1 < len(parts) and re.match(r'^([A-Z]\.[\s-]?)+$', parts[i+1].strip()):
            authors.append((part, parts[i+1].strip()))
            i += 2
        elif ',' in part:
            comma_idx = part.index(',')
            last = part[:comma_idx].strip()
            inits = part[comma_idx+1:].strip()
            authors.append((last, inits))
            i += 1
        else:
            authors.append((part, ''))
            i += 1
    return authors

def format_authors_apa7(authors):
    if not authors: return ''
    formatted = []
    for last, inits in authors:
        inits_clean = re.sub(r'([A-Z])(?!\.)(?!\s*[A-Z]\.)', r'\1.', inits)
        inits_clean = re.sub(r'\s+', ' ', inits_clean).strip()
        if inits_clean:
            formatted.append(f"{last}, {inits_clean}")
        else:
            formatted.append(last)
    if len(formatted) == 1:
        return formatted[0]
    elif len(formatted) <= 20:
        return ', '.join(formatted[:-1]) + ', & ' + formatted[-1]
    else:
        return ', '.join(formatted[:19]) + ', . . . ' + formatted[-1]

def detect_ref_type(ref_text):
    t = ref_text.lower()
    if 'http' in t or 'doi' in t or 'retrieved from' in t or 'diakses' in t:
        if any(x in t for x in ['skripsi','thesis','tesis','dissertation','disertasi']):
            return 'thesis_online'
        if any(x in t for x in ['journal','jurnal']) or re.search(r'\d+\(\d+\)', t):
            return 'journal'
        return 'webpage'
    if any(x in t for x in ['skripsi','thesis','tesis','dissertation','disertasi']):
        return 'thesis'
    if re.search(r',\s*\d+\s*\(\d+\)', ref_text):
        return 'journal'
    if re.search(r'\.\s+[A-Z][^.]+\.$', ref_text.strip()):
        return 'book'
    return 'journal'

def parse_reference(ref_text):
    ref_text = clean_text(ref_text)
    result = {
        'raw': ref_text, 'authors': '', 'year': '', 'title': '',
        'source': '', 'volume': '', 'issue': '', 'pages': '',
        'doi': '', 'publisher': '', 'ref_type': 'journal', 'url': '', 'thesis_info': '',
    }
    doi_match = re.search(r'https?://doi\.org/\S+|doi:\s*\S+', ref_text, re.IGNORECASE)
    if doi_match:
        result['doi'] = doi_match.group().strip().rstrip('.')
        ref_text = ref_text[:doi_match.start()].strip()
    url_match = re.search(r'https?://\S+', ref_text)
    if url_match and not result['doi']:
        result['url'] = url_match.group().strip().rstrip('.')
        ref_text = ref_text[:url_match.start()].strip()
    year_match = re.search(r'\((\d{4}[a-z]?)\)', ref_text)
    if year_match:
        result['year'] = year_match.group(1)
        authors_part = ref_text[:year_match.start()].strip().rstrip('.,')
        rest = ref_text[year_match.end():].strip().lstrip('.,').strip()
    else:
        authors_part = ref_text; rest = ''
    result['authors'] = authors_part
    ref_type = detect_ref_type(ref_text)
    result['ref_type'] = ref_type
    if ref_type == 'journal':
        parts = [p.strip() for p in rest.split('.') if p.strip()]
        if len(parts) >= 1: result['title'] = parts[0]
        if len(parts) >= 2:
            jp = parts[1]
            vm = re.search(r',\s*(\d+)\s*\((\d+)\)\s*,\s*([\d\u2013\-]+)', jp)
            if vm:
                result['source'] = jp[:vm.start()].strip()
                result['volume'] = vm.group(1); result['issue'] = vm.group(2); result['pages'] = vm.group(3)
            else:
                vm2 = re.search(r',\s*(\d+)\s*,\s*([\d\u2013\-]+)', jp)
                if vm2:
                    result['source'] = jp[:vm2.start()].strip()
                    result['volume'] = vm2.group(1); result['pages'] = vm2.group(2)
                else:
                    result['source'] = jp
    elif ref_type == 'book':
        parts = [p.strip() for p in rest.split('.') if p.strip()]
        if parts: result['title'] = parts[0]
        if len(parts) >= 2: result['publisher'] = parts[1]
    elif ref_type in ('thesis', 'thesis_online'):
        parts = [p.strip() for p in rest.split('.') if p.strip()]
        if parts: result['title'] = parts[0]
        tm = re.search(r'\(([^)]+(?:skripsi|thesis|tesis|dissertation|disertasi)[^)]*)\)', rest, re.IGNORECASE)
        if tm: result['thesis_info'] = tm.group(1)
        im = re.search(r'(?:Universitas|University|Institut|Institute)[^.,)]+', rest, re.IGNORECASE)
        if im: result['source'] = im.group().strip()
    elif ref_type == 'webpage':
        parts = [p.strip() for p in rest.split('.') if p.strip()]
        if parts: result['title'] = parts[0]
        if len(parts) >= 2: result['source'] = parts[1]
    return result

def format_reference_apa7(parsed):
    r = parsed
    try:
        al = parse_author_string(r['authors'])
        authors_str = format_authors_apa7(al)
    except:
        authors_str = r['authors']
    year = f"({r['year']})" if r['year'] else ''
    segs = []
    if r['ref_type'] == 'journal':
        vi = f"{r['volume']}({r['issue']})" if r['volume'] and r['issue'] else r['volume']
        pages = r['pages'].strip().rstrip('.') if r['pages'] else ''
        segs.append((f"{authors_str} {year}. {r['title'].strip().rstrip('.')}. ", False))
        jp = r['source'].strip().rstrip('.')
        if vi: jp += f", {vi}"
        if pages: jp += f", {pages}"
        segs.append((jp, True))
        if r['doi']: segs.append((f". {r['doi']}", False))
        elif r['url']: segs.append((f". {r['url']}", False))
        else: segs.append((".", False))
    elif r['ref_type'] == 'book':
        segs.append((f"{authors_str} {year}. ", False))
        segs.append((f"{r['title'].strip().rstrip('.')}. ", True))
        segs.append((f"{r['publisher'].strip().rstrip('.')}.", False))
    elif r['ref_type'] in ('thesis', 'thesis_online'):
        ti = r['thesis_info'] if r['thesis_info'] else 'Skripsi'
        src = r['source'].strip().rstrip('.') if r['source'] else ''
        segs.append((f"{authors_str} {year}. ", False))
        segs.append((r['title'].strip().rstrip('.'), True))
        if src: segs.append((f" ({ti}, {src}).", False))
        else: segs.append((f" ({ti}).", False))
        if r['url']: segs.append((f" Retrieved from {r['url']}", False))
        if r['doi']: segs.append((f" {r['doi']}", False))
    elif r['ref_type'] == 'webpage':
        segs.append((f"{authors_str} {year}. ", False))
        segs.append((f"{r['title'].strip().rstrip('.')}. ", True))
        if r['source']: segs.append((f"{r['source'].strip().rstrip('.')}.", False))
        if r['url']: segs.append((f" {r['url']}", False))
        if r['doi']: segs.append((f" {r['doi']}", False))
    else:
        segs.append((r['raw'], False))
    return segs

# ─────────────────────────────────────────────
# DOCUMENT FORMATTER
# ─────────────────────────────────────────────

def set_page_margins(doc):
    for section in doc.sections:
        section.page_width = Cm(21.0); section.page_height = Cm(29.7)
        section.top_margin = Cm(3.0); section.bottom_margin = Cm(3.0)
        section.left_margin = Cm(3.0); section.right_margin = Cm(3.0)

def clear_pf(para):
    pf = para.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    pf.first_line_indent = None; pf.left_indent = None; pf.right_indent = None

def set_run_font(run, fname, fsize, bold=False, italic=False):
    run.font.name = fname; run.font.size = Pt(fsize)
    run.font.bold = bold; run.font.italic = italic
    rpr = run._r.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rpr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), fname); rFonts.set(qn('w:hAnsi'), fname); rFonts.set(qn('w:cs'), fname)

def get_full_text(para):
    return ''.join(r.text for r in para.runs).strip()

def classify_paragraph(para, prev_text=''):
    style_name = (para.style.name or '').lower()
    text = get_full_text(para)
    text_lower = text.lower().strip()
    if not text_lower: return 'empty'
    if 'heading 1' in style_name: return 'h1'
    if 'heading 2' in style_name: return 'h2'
    if 'heading 3' in style_name: return 'h3'
    if 'title' in style_name: return 'title'
    if re.match(r'^(abstract|abstrak)\s*$', text_lower): return 'abstract_label'
    if re.match(r'^(kata kunci|keywords?)\s*:', text_lower): return 'keywords'
    if re.match(r'^received|^diterima', text_lower): return 'received_line'
    if text_lower in SECTIONS_H1: return 'h1'
    if re.match(r'^[A-Z][a-zA-Z\-]+,\s+[A-Z]\.', text): return 'reference_entry'
    if re.match(r'^(table|tabel)\s+\d+', text_lower): return 'table_title'
    if re.match(r'^(figure|gambar)\s+\d+', text_lower): return 'figure_caption'
    if re.match(r'^email:', text_lower) or ('@' in text and len(text) < 80): return 'email'
    all_bold = all(r.bold for r in para.runs if r.text.strip())
    if all_bold and len(text) < 80:
        if text_lower in SECTIONS_H1: return 'h1'
        if text_lower in SECTIONS_H2: return 'h2'
        if len(text) < 50: return 'h2'
    return 'body'

def fmt_title(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_after = Pt(20)
    for r in para.runs: set_run_font(r, FONT_ARIAL, 14, bold=True)

def fmt_author(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in para.runs: set_run_font(r, FONT_TNR, 10, bold=True)

def fmt_affil(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in para.runs: set_run_font(r, FONT_TNR, 10)

def fmt_email(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_after = Pt(10)
    for r in para.runs: set_run_font(r, FONT_TNR, 10)

def fmt_abstract_label(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(10)
    for r in para.runs: set_run_font(r, FONT_TNR, 10, bold=True)

def fmt_abstract_body(para, en=False):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.first_line_indent = Cm(0)
    for r in para.runs: set_run_font(r, FONT_TNR, 10, italic=en)

def fmt_keywords(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in para.runs: set_run_font(r, FONT_TNR, 10, bold=True, italic=True)

def fmt_received(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(10); para.paragraph_format.space_after = Pt(20)
    for r in para.runs: set_run_font(r, FONT_TNR, 10, italic=True)

def fmt_h1(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(10)
    for r in para.runs: set_run_font(r, FONT_ARIAL, 12, bold=True)

def fmt_h2(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(10)
    for r in para.runs: set_run_font(r, FONT_ARIAL, 10, bold=True)

def fmt_h3(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(10)
    for r in para.runs: set_run_font(r, FONT_ARIAL, 10, italic=True)

def fmt_body(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.first_line_indent = Cm(0.9)
    for r in para.runs: set_run_font(r, FONT_TNR, 10)

def fmt_table_title(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(10); para.paragraph_format.space_after = Pt(2)
    for r in para.runs: set_run_font(r, FONT_TNR, 10, bold=True)

def fmt_fig_caption(para):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(2); para.paragraph_format.space_after = Pt(10)
    for r in para.runs: set_run_font(r, FONT_TNR, 10, italic=True)

def fmt_ref_entry(para, segments):
    clear_pf(para); para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.left_indent = Cm(0.9)
    para.paragraph_format.first_line_indent = Cm(-0.9)
    if segments:
        for r in para.runs: r.text = ''
        first = True
        for text, is_italic in segments:
            if first and para.runs:
                r = para.runs[0]; r.text = text; first = False
            else:
                r = para.add_run(text)
            set_run_font(r, FONT_TNR, 10, italic=is_italic)
    else:
        for r in para.runs: set_run_font(r, FONT_TNR, 10)

def fmt_table(table):
    n_rows = len(table.rows)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in para.runs:
                    set_run_font(r, FONT_TNR, 10, bold=(row_idx == 0))
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr'); tbl.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement(f'w:{side}')
        if side in ('top','bottom'):
            b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'6'); b.set(qn('w:color'),'000000')
        else:
            b.set(qn('w:val'),'none'); b.set(qn('w:sz'),'0'); b.set(qn('w:color'),'auto')
        tblBorders.append(b)
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None: tblPr.remove(old)
    tblPr.append(tblBorders)
    if n_rows > 1:
        for cell in table.rows[0].cells:
            tc = cell._tc
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is None: tcPr = OxmlElement('w:tcPr'); tc.insert(0, tcPr)
            tcBorders = OxmlElement('w:tcBorders')
            b = OxmlElement('w:bottom')
            b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'6'); b.set(qn('w:color'),'000000')
            tcBorders.append(b)
            old_tb = tcPr.find(qn('w:tcBorders'))
            if old_tb is not None: tcPr.remove(old_tb)
            tcPr.append(tcBorders)

def format_document(input_bytes):
    """Format document from bytes, return formatted bytes."""
    doc = Document(io.BytesIO(input_bytes))
    set_page_margins(doc)
    paragraphs = list(doc.paragraphs)
    classifications = []
    prev_text = ''
    for para in paragraphs:
        cls = classify_paragraph(para, prev_text)
        classifications.append(cls)
        prev_text = get_full_text(para)

    state = 'preamble'
    id_abstract_active = False
    en_abstract_active = False
    ref_section_active = False

    for i, para in enumerate(paragraphs):
        text = get_full_text(para)
        text_lower = text.lower().strip()
        cls = classifications[i]

        if text_lower == 'abstrak':
            state = 'abstract_id_label'; id_abstract_active = True; en_abstract_active = False
        elif text_lower == 'abstract':
            if state in ('abstract_id_label','abstract_id_body','keywords_id'):
                state = 'abstract_en_label'; id_abstract_active = False; en_abstract_active = True
            else:
                state = 'body'; en_abstract_active = False
        elif re.match(r'^(kata kunci|keywords?)\s*:', text_lower):
            if id_abstract_active: state = 'keywords_id'
            elif en_abstract_active: state = 'keywords_en'; en_abstract_active = False
        elif re.match(r'^received|^diterima', text_lower):
            state = 'received'; en_abstract_active = False; id_abstract_active = False
        elif text_lower in SECTIONS_H1:
            if 'reference' in text_lower or 'pustaka' in text_lower:
                ref_section_active = True; state = 'references'
            else:
                state = 'body'; ref_section_active = False
        elif ref_section_active and re.match(r'^[A-Z][a-zA-Z\-]+,?\s+[A-Z]', text):
            cls = 'reference_entry'; classifications[i] = cls

        if cls == 'empty':
            clear_pf(para); para.paragraph_format.space_before = Pt(0); para.paragraph_format.space_after = Pt(0)
            continue

        if state == 'abstract_id_label' and text_lower == 'abstrak': fmt_abstract_label(para)
        elif state == 'abstract_en_label' and text_lower == 'abstract': fmt_abstract_label(para)
        elif id_abstract_active and text_lower != 'abstrak' and cls not in ('keywords',):
            state = 'abstract_id_body'; fmt_abstract_body(para, en=False)
        elif en_abstract_active and text_lower != 'abstract' and cls not in ('keywords',):
            state = 'abstract_en_body'; fmt_abstract_body(para, en=True)
        elif state == 'keywords_id': fmt_keywords(para)
        elif state == 'keywords_en': fmt_keywords(para)
        elif state == 'received': fmt_received(para); state = 'body'
        elif cls == 'h1': fmt_h1(para)
        elif cls == 'h2': fmt_h2(para)
        elif cls == 'h3': fmt_h3(para)
        elif cls == 'table_title': fmt_table_title(para)
        elif cls == 'figure_caption': fmt_fig_caption(para)
        elif cls == 'reference_entry' or (ref_section_active and re.match(r'^[A-Z]', text) and len(text) > 30):
            parsed = parse_reference(text)
            segs = format_reference_apa7(parsed)
            fmt_ref_entry(para, segs)
        elif cls == 'title' or (i < 4 and len(text) > 20): fmt_title(para)
        elif cls in ('author',): fmt_author(para)
        elif cls in ('affiliation',): fmt_affil(para)
        elif cls == 'email': fmt_email(para)
        else: fmt_body(para)

    for table in doc.tables:
        fmt_table(table)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()

# ─────────────────────────────────────────────
# STREAMLIT UI
# ─────────────────────────────────────────────

st.set_page_config(
    page_title="JIKAP Manuscript Formatter",
    page_icon="📄",
    layout="centered"
)

# Header
st.title("📄 JIKAP Manuscript Formatter")
st.markdown(
    "*Jurnal Informasi dan Komunikasi Administrasi Perkantoran · PAP FKIP UNS*"
)
st.divider()

# Info box
with st.expander("ℹ️ Format yang diterapkan", expanded=False):
    st.markdown("""
    | Elemen | Format |
    |---|---|
    | Halaman | A4, margin 3 cm semua sisi |
    | Judul artikel | Arial 14pt Bold |
    | Body teks | Times New Roman 10pt, justified, indent |
    | Heading 1 | Arial 12pt Bold |
    | Heading 2 | Arial 10pt Bold |
    | Heading 3 | Arial 10pt Italic |
    | Abstrak Indonesia | TNR 10pt tegak |
    | Abstrak Inggris | TNR 10pt italic |
    | Keywords | TNR 10pt Bold Italic |
    | Referensi | TNR 10pt, hanging indent, APA 7 |
    | Tabel | Border atas/bawah header/bawah; judul di atas |
    | Caption gambar | TNR 10pt italic, di bawah gambar |
    """)

# Upload
st.subheader("1. Upload Manuscript")
uploaded_file = st.file_uploader(
    "Pilih file manuscript (.docx)",
    type=["docx"],
    help="File harus dalam format Microsoft Word (.docx)"
)

if uploaded_file is not None:
    st.success(f"✅ File diterima: **{uploaded_file.name}** ({uploaded_file.size / 1024:.1f} KB)")

    st.subheader("2. Format Manuscript")

    if st.button("▶ Format Sekarang", type="primary", use_container_width=True):
        with st.spinner("Memformat manuscript..."):
            try:
                input_bytes = uploaded_file.getvalue()
                output_bytes = format_document(input_bytes)

                st.success("✅ Manuscript berhasil diformat!")
                st.subheader("3. Download Hasil")

                output_filename = uploaded_file.name.replace(".docx", "_JIKAP_formatted.docx")
                st.download_button(
                    label="⬇️ Download Manuscript Terformat",
                    data=output_bytes,
                    file_name=output_filename,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True,
                    type="primary"
                )

            except Exception as e:
                st.error(f"❌ Terjadi error: {str(e)}")
                st.info("Pastikan file adalah dokumen Word (.docx) yang valid.")

else:
    st.info("👆 Upload file manuscript untuk memulai.")

# Footer
st.divider()
st.caption("JIKAP Manuscript Formatter · PAP FKIP Universitas Sebelas Maret · 2025")
